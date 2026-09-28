"""
Convierte los JSON que genera fetch_benefits.py en un Excel.

    python fetch_benefits.py          # 1) descarga (crea output/*.json)
    python to_excel.py                # 2) crea output/beneficios.xlsx

Hojas: "Todos" (todos los bancos juntos, con columna `archivo`) y una hoja por archivo JSON.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

COLUMNS = ["source", "title", "discount", "days", "category", "location", "valid_from", "valid_to",
           "description", "conditions", "url", "merchant", "tags", "images", "id"]
WIDTHS = {"title": 35, "description": 60, "conditions": 60, "url": 45, "location": 35, "tags": 30, "images": 45}
EXCEL_CELL_LIMIT = 32767


def cell(value):
    if isinstance(value, list):
        value = " | ".join(str(v) for v in value)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)  # caracteres de control que Excel rechaza
    if isinstance(value, str) and len(value) > EXCEL_CELL_LIMIT:
        value = value[:EXCEL_CELL_LIMIT - 1] + "…"
    return value


def add_sheet(wb: Workbook, title: str, rows: list[dict], extra_first: str | None = None):
    ws = wb.create_sheet(title[:31])  # Excel limita los nombres de hoja a 31 caracteres
    headers = ([extra_first] if extra_first else []) + COLUMNS
    ws.append(headers)
    for row in rows:
        ws.append([cell(row.get(h, "")) for h in headers])
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(h, 16)
        ws.cell(1, i).font = Font(bold=True)
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(vertical="top", wrap_text=False)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--src", type=Path, default=Path("./output"), help="carpeta con los JSON")
    ap.add_argument("--out", type=Path, default=None, help="archivo .xlsx (por defecto <src>/beneficios.xlsx)")
    args = ap.parse_args()
    out = args.out or args.src / "beneficios.xlsx"

    files = sorted(p for p in args.src.glob("*.json") if not p.name.startswith("_"))
    if not files:
        raise SystemExit(f"No hay archivos JSON en {args.src.resolve()}. Corre primero: python fetch_benefits.py")

    wb = Workbook()
    wb.remove(wb.active)
    everything, per_file = [], []
    for f in files:
        rows = json.loads(f.read_text(encoding="utf-8"))
        per_file.append((f.stem, rows))
        everything += [{"archivo": f.stem, **r} for r in rows]
        print(f"{f.name:<40} {len(rows):>5} registros")

    add_sheet(wb, "Todos", everything, extra_first="archivo")
    for name, rows in per_file:
        add_sheet(wb, name, rows)
    wb.save(out)
    print(f"\n{len(everything)} registros -> {out.resolve()}")


if __name__ == "__main__":
    main()
